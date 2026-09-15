import openpyxl
from openpyxl.utils import get_column_letter

path = r"Y-Global2 2017-20 (Use this one).xlsx"
wb = openpyxl.load_workbook(path, data_only=True)


def row_values(ws, r, c1=1, c2=52):
    out = {}
    for c in range(c1, c2 + 1):
        v = ws.cell(r, c).value
        if v is not None and str(v).strip() != "":
            out[get_column_letter(c)] = v
    return out

for sh in ["Oper Flow 25 ", "Oper Flow 26 "]:
    ws = wb[sh]
    print(f"\n===== {sh} =====")
    for r in [2,3,10,13,34,35,36,37,38,39,40]:
        vals = row_values(ws, r, 1, 53)
        if vals:
            pretty = " | ".join([f"{k}={v}" for k,v in vals.items()])
            print(f"R{r}: {pretty}")

# Estado financiero: extraer saldos apertura y cierre aproximado por mes clave
for sh in ["ENE25", "DIC25", "ENE26", "FEB26", "MAR26", "ABR26"]:
    if sh not in wb.sheetnames:
        continue
    ws = wb[sh]
    print(f"\n===== {sh} =====")

    # Buscar filas con etiquetas relevantes
    keywords = ["Saldo Apertura", "Saldo Cierre", "Cierre", "Balance", "Comprobacion", "TOTAL"]
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        tag = f"{a} {b}".strip()
        if any(k.lower() in tag.lower() for k in keywords if isinstance(tag, str)):
            vals = row_values(ws, r, 1, 16)
            if vals:
                print(f"R{r}: " + " | ".join([f"{k}={v}" for k,v in vals.items()]))

    # imprimir ultima fila con datos en C:P
    last = None
    for r in range(ws.max_row, 0, -1):
        has = False
        for c in range(3, 17):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip() != "":
                has = True
                break
        if has:
            last = r
            break
    if last:
        vals = row_values(ws, last, 1, 16)
        print(f"LAST R{last}: " + " | ".join([f"{k}={v}" for k,v in vals.items()]))
