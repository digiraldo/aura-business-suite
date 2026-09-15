import csv
import openpyxl
from datetime import date

xlsx = r"Y-Global2 2017-20 (Use this one).xlsx"
out_csv = r"import_transacciones_2026_completo.csv"

wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb["Oper Flow 26 "]

months = [
    ("ENE", 1, 11, 12),  # (name, month_num, presupuesto_col, actual_col)
    ("FEB", 2, 13, 14),
    ("MAR", 3, 15, 16),
    ("ABR", 4, 17, 18),
    ("MAY", 5, 19, 20),
    ("JUN", 6, 21, 22),
    ("JUL", 7, 23, 24),
    ("AGO", 8, 25, 26),
    ("SEP", 9, 27, 28),
    ("OCT", 10, 29, 30),
    ("NOV", 11, 31, 32),
    ("DIC", 12, 33, 34),
]

def num(v):
    try:
        return float(v)
    except Exception:
        return 0.0

rows = []

# Ingresos: filas 4-9
for r in range(4, 10):
    cat = ws.cell(r, 1).value
    if not cat:
        continue
    for mname, mnum, c_presu, c_act in months:
        presu = num(ws.cell(r, c_presu).value)
        act = num(ws.cell(r, c_act).value)
        if act > 0:
            amount = act
            src = "actual"
        else:
            amount = presu
            src = "presupuesto"
        if amount <= 0:
            continue

        d = date(2026, mnum, 1).strftime("%d/%m/%Y")
        desc = f"[{mname}26][income][{src}] {cat}"
        note = f"Origen: Oper Flow 26, fila {r}. Valor {src}."
        ref = f"{mname}26-INC-R{r}"
        rows.append([d, "income", str(cat), f"{amount:.2f}", desc, note, "transferencia", ref])

# Egresos: filas 14-32
for r in range(14, 33):
    cat = ws.cell(r, 1).value
    if not cat:
        continue
    for mname, mnum, c_presu, c_act in months:
        presu = num(ws.cell(r, c_presu).value)
        act = num(ws.cell(r, c_act).value)
        if act > 0:
            amount = act
            src = "actual"
        else:
            amount = presu
            src = "presupuesto"
        if amount <= 0:
            continue

        d = date(2026, mnum, 1).strftime("%d/%m/%Y")
        desc = f"[{mname}26][expense][{src}] {cat}"
        note = f"Origen: Oper Flow 26, fila {r}. Valor {src}."
        ref = f"{mname}26-EXP-R{r}"
        rows.append([d, "expense", str(cat), f"{amount:.2f}", desc, note, "transferencia", ref])

# ordenar por fecha y tipo
rows.sort(key=lambda x: (x[0].split('/')[2], x[0].split('/')[1], x[0].split('/')[0], x[1], x[2]))

with open(out_csv, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["transaction_date", "transaction_type", "category", "amount", "description", "notes", "payment_method", "reference_number"])
    w.writerows(rows)

print(f"rows={len(rows)} file={out_csv}")
