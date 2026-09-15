import json
import openpyxl

path = r"Y-Global2 2017-20 (Use this one).xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
print(json.dumps({"sheets": wb.sheetnames}, ensure_ascii=False))
for s in wb.sheetnames:
    ws = wb[s]
    print(f"{s}\trows={ws.max_row}\tcols={ws.max_column}")
