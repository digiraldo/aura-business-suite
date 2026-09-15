import openpyxl

wb = openpyxl.load_workbook(r"Y-Global2 2017-20 (Use this one).xlsx", data_only=True)
ws = wb["Oper Flow 26 "]

# 2025 actual por rubro está en columna I
INCOME_START, INCOME_END = 4, 9
EXP_START, EXP_END = 14, 32

inc = []
for r in range(INCOME_START, INCOME_END+1):
    name = ws.cell(r,1).value
    val = ws.cell(r,9).value
    if isinstance(val,(int,float)):
        inc.append((name,float(val)))

exp = []
for r in range(EXP_START, EXP_END+1):
    name = ws.cell(r,1).value
    val = ws.cell(r,9).value
    if isinstance(val,(int,float)):
        exp.append((name,float(val)))

inc.sort(key=lambda x: x[1], reverse=True)
exp.sort(key=lambda x: x[1], reverse=True)

print("TOP INGRESOS 2025")
for n,v in inc:
    print(f"{n}\t{v:.2f}")

print("\nTOP EGRESOS 2025")
for n,v in exp:
    print(f"{n}\t{v:.2f}")
